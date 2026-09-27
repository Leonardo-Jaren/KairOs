"""Renderiza el encabezado corporativo y los indicadores KPI del reporte."""

from datetime import datetime
from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .excel_theme import ExcelTheme


class ExcelHeaderRenderer:
    def __init__(self, worksheet, title: str, user_name: str, filters_text: str):
        self.worksheet = worksheet
        self.title = title
        self.user_name = user_name
        self.filters_text = filters_text

    def render_header(self) -> int:
        worksheet = self.worksheet

        worksheet.merge_cells("A1:K1")
        banner = worksheet["A1"]
        banner.value = "  KAIROS · GESTIÓN INTEGRAL DE INFRAESTRUCTURA Y SOPORTE"
        banner.font = Font(name="Calibri", size=10, bold=True, color=ExcelTheme.COLOR_WHITE)
        banner.fill = PatternFill(
            start_color=ExcelTheme.COLOR_NAVY_DARK,
            end_color=ExcelTheme.COLOR_NAVY_DARK,
            fill_type="solid",
        )
        banner.alignment = Alignment(vertical="center", horizontal="left")
        worksheet.row_dimensions[1].height = 24

        worksheet.merge_cells("A2:K2")
        report_title = worksheet["A2"]
        report_title.value = f"  {self.title.upper()}"
        report_title.font = Font(name="Calibri", size=15, bold=True, color=ExcelTheme.COLOR_NAVY_DARK)
        report_title.alignment = Alignment(vertical="center", horizontal="left")
        worksheet.row_dimensions[2].height = 30

        emitted_at = datetime.now().strftime("%d/%m/%Y %H:%M")
        metadata = f"  Generado por: {self.user_name or 'Administrador'}  |  Fecha de emisión: {emitted_at}"
        if self.filters_text:
            metadata += f"  |  Filtros aplicados: {self.filters_text}"
        else:
            metadata += "  |  Alcance: Todos los registros (Sin filtros específicos)"

        worksheet.merge_cells("A3:K3")
        metadata_cell = worksheet["A3"]
        metadata_cell.value = metadata
        metadata_cell.font = Font(name="Calibri", size=9, italic=True, color=ExcelTheme.COLOR_TEXT_MUTED)
        metadata_cell.alignment = Alignment(vertical="center", horizontal="left")
        worksheet.row_dimensions[3].height = 20
        worksheet.row_dimensions[4].height = 8
        return 5

    def add_kpis(self, current_row: int, kpi_list: list[dict[str, Any]]) -> int:
        if not kpi_list:
            return current_row

        worksheet = self.worksheet
        worksheet.row_dimensions[current_row].height = 18
        worksheet.row_dimensions[current_row + 1].height = 26

        side = Side(border_style="thin", color=ExcelTheme.COLOR_BORDER)
        kpi_border = Border(top=side, left=side, right=side, bottom=side)

        for index, kpi in enumerate(kpi_list):
            first_column = index * 2 + 1
            label_cell = worksheet.cell(row=current_row, column=first_column)
            value_cell = worksheet.cell(row=current_row + 1, column=first_column)
            end_column = get_column_letter(first_column + 1)
            start_column = get_column_letter(first_column)
            worksheet.merge_cells(f"{start_column}{current_row}:{end_column}{current_row}")
            worksheet.merge_cells(f"{start_column}{current_row + 1}:{end_column}{current_row + 1}")

            palette = ExcelTheme.STATUS_PALETTES.get(
                kpi.get("tone", "neutral"), ExcelTheme.STATUS_PALETTES['neutral']
            )
            label_cell.value = kpi.get("label", "").upper()
            label_cell.font = Font(name="Calibri", size=8, bold=True, color=ExcelTheme.COLOR_TEXT_MUTED)
            label_cell.fill = PatternFill(start_color=ExcelTheme.COLOR_PRIMARY_BG, fill_type="solid")
            label_cell.alignment = Alignment(horizontal="center", vertical="center")
            label_cell.border = kpi_border

            value_cell.value = kpi.get("value", 0)
            value_cell.font = Font(name="Calibri", size=14, bold=True, color=palette['font'])
            value_cell.fill = PatternFill(start_color=palette['fill'], fill_type="solid")
            value_cell.alignment = Alignment(horizontal="center", vertical="center")
            value_cell.border = kpi_border

        next_row = current_row + 3
        worksheet.row_dimensions[next_row - 1].height = 10
        return next_row
