"""Pruebas del contrato público del constructor Excel compartido."""

from io import BytesIO

from django.test import SimpleTestCase
from openpyxl import load_workbook

from shared.reports import ExcelReportBuilder


class ExcelReportBuilderTests(SimpleTestCase):
    def test_renderiza_kpis_tabla_estados_y_configuracion_de_lectura(self):
        builder = ExcelReportBuilder(
            title="Reporte de prueba",
            user_name="Agente de prueba",
            filters_text="Estado: en uso",
            sheet_name="Equipos",
        )
        builder.add_kpis([{"label": "Total", "value": 1, "tone": "info"}])
        builder.set_columns([
            {"key": "codigo", "label": "Código", "width": 18},
            {"key": "estado", "label": "Estado", "width": 16},
        ])
        builder.add_rows([{"codigo": "EQ-001", "estado": "en_uso"}], status_keys=['estado'])

        workbook = load_workbook(BytesIO(builder.to_bytes()))
        worksheet = workbook["Equipos"]

        self.assertEqual(worksheet["A2"].value, "  REPORTE DE PRUEBA")
        self.assertIn("Agente de prueba", worksheet["A3"].value)
        self.assertEqual(worksheet["A5"].value, "TOTAL")
        self.assertEqual(worksheet["A8"].value, "Código")
        self.assertEqual(worksheet["A9"].value, "EQ-001")
        self.assertEqual(worksheet.freeze_panes, "A9")
        self.assertEqual(worksheet.auto_filter.ref, "A8:B9")
        self.assertEqual(worksheet["B9"].fill.fgColor.rgb, "00D1FAE5")

    def test_add_sheet_reinicia_la_tabla_y_conserva_el_contrato_de_respuesta(self):
        builder = ExcelReportBuilder(title="Reporte", sheet_name="Resumen")
        builder.set_columns([{"key": "x", "label": "X"}]).add_rows([{"x": 1}])
        builder.add_sheet("Detalle")
        builder.set_columns([{"key": "valor", "label": "Valor"}]).add_rows([{"valor": 2}])

        response = builder.to_response("detalle")
        workbook = load_workbook(BytesIO(response.content))

        self.assertEqual(workbook.sheetnames, ["Resumen", "Detalle"])
        self.assertEqual(workbook["Detalle"]["A5"].value, "Valor")
        self.assertEqual(workbook["Detalle"]["A6"].value, 2)
        self.assertEqual(response["Content-Disposition"], 'attachment; filename="detalle.xlsx"')
        self.assertEqual(response["Access-Control-Expose-Headers"], "Content-Disposition")
