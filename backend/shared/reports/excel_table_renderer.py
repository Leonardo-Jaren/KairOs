"""Renderiza tablas de reportes: columnas, filas y configuración de lectura."""

from datetime import date, datetime
from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .excel_theme import ExcelTheme


class ExcelTableRenderer:
    def __init__(self, worksheet):
        self.worksheet = worksheet
        self.columns: list[dict[str, Any]] = []
        self.header_row = 1
        self.data_start_row = 1

    def set_worksheet(self, worksheet) -> None:
        self.worksheet = worksheet
        self.columns = []
        self.header_row = 1
        self.data_start_row = 1

    def set_columns(self, columns: list[dict[str, Any]], current_row: int) -> tuple[int, int]:
        self.columns = columns
        self.header_row = current_row
        self.data_start_row = current_row + 1
        worksheet = self.worksheet
        worksheet.row_dimensions[self.header_row].height = 26

        header_fill = PatternFill(start_color=ExcelTheme.COLOR_NAVY_LIGHT, fill_type="solid")
        header_font = Font(name="Calibri", size=10, bold=True, color=ExcelTheme.COLOR_WHITE)
        header_border = Border(
            top=Side(style="thin", color=ExcelTheme.COLOR_BORDER),
            left=Side(style="thin", color=ExcelTheme.COLOR_BORDER),
            right=Side(style="thin", color=ExcelTheme.COLOR_BORDER),
            bottom=Side(style="medium", color=ExcelTheme.COLOR_NAVY_DARK),
        )

        for column_index, column in enumerate(columns, start=1):
            cell = worksheet.cell(row=self.header_row, column=column_index, value=column.get("label", ""))
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal=column.get("align", "left"),
                vertical="center",
                wrap_text=True,
            )
            cell.border = header_border
            if "width" in column:
                worksheet.column_dimensions[get_column_letter(column_index)].width = column["width"]

        return self.header_row, self.data_start_row

    def add_rows(
        self,
        data: list[dict[str, Any]],
        current_row: int,
        status_keys: list[str] | None = None,
    ) -> int:
        if not self.columns:
            raise ValueError("Debe definir las columnas con set_columns antes de agregar filas.")

        status_keys = status_keys or ['estado', 'prioridad', 'resultado_equipo', 'is_active']
        worksheet = self.worksheet
        side = Side(style="thin", color=ExcelTheme.COLOR_BORDER)
        cell_border = Border(top=side, left=side, right=side, bottom=side)
        data_font = Font(name="Calibri", size=10, color=ExcelTheme.COLOR_TEXT_MAIN)
        mono_font = Font(name="Consolas", size=9, color=ExcelTheme.COLOR_TEXT_MAIN)

        for row_index, item in enumerate(data, start=current_row):
            worksheet.row_dimensions[row_index].height = 20
            row_fill_color = ExcelTheme.COLOR_ZEBRA if row_index % 2 == 1 else ExcelTheme.COLOR_WHITE
            row_fill = PatternFill(start_color=row_fill_color, fill_type="solid")

            for column_index, column in enumerate(self.columns, start=1):
                key = column["key"]
                raw_value = item.get(key)
                alignment = column.get("align", "left")
                value = raw_value

                if isinstance(raw_value, (datetime, date)):
                    value = raw_value.strftime("%d/%m/%Y")
                    alignment = "center"
                elif raw_value is None:
                    value = "—"
                    alignment = "center"
                elif isinstance(raw_value, bool):
                    value = "Sí" if raw_value else "No"
                    alignment = "center"

                cell = worksheet.cell(row=row_index, column=column_index, value=value)
                cell.font = mono_font if column.get("mono", False) else data_font
                cell.border = cell_border
                cell.alignment = Alignment(horizontal=alignment, vertical="center")
                cell.fill = row_fill

                if key in status_keys and raw_value is not None:
                    normalized_status = str(raw_value).lower().strip()
                    tone = ExcelTheme.STATUS_MAP.get(normalized_status)
                    if tone:
                        palette = ExcelTheme.STATUS_PALETTES[tone]
                        cell.fill = PatternFill(start_color=palette['fill'], fill_type="solid")
                        cell.font = Font(name="Calibri", size=9, bold=True, color=palette['font'])
                        cell.alignment = Alignment(horizontal="center", vertical="center")

        next_row = current_row + len(data)
        self._finalize_sheet_formatting(next_row)
        return next_row

    def _finalize_sheet_formatting(self, current_row: int) -> None:
        worksheet = self.worksheet
        if not self.columns:
            return

        last_row = max(self.data_start_row, current_row - 1)
        last_column = get_column_letter(len(self.columns))
        worksheet.freeze_panes = f"A{self.data_start_row}"
        worksheet.auto_filter.ref = f"A{self.header_row}:{last_column}{last_row}"

        for column_index, column in enumerate(self.columns, start=1):
            if "width" in column:
                continue

            column_letter = get_column_letter(column_index)
            max_length = len(str(column.get("label", ""))) + 4
            sample_step = max(1, (last_row - self.data_start_row) // 50)
            for row_index in range(self.data_start_row, last_row + 1, sample_step):
                value = worksheet.cell(row=row_index, column=column_index).value
                if value:
                    max_length = max(max_length, len(str(value)) + 3)
            worksheet.column_dimensions[column_letter].width = min(max(max_length, 12), 48)
