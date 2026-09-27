"""Renderiza el reporte técnico multihoja de un espacio y su croquis."""

from datetime import datetime

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from shared.reports import ExcelReportBuilder


class PlanoEspacioExcelReport:
    @staticmethod
    def generar(instance, actor=None):
        actor_name = f"{actor.nombre} {actor.apellido}".strip() if actor else "Administrador"
        equipos = getattr(instance, 'equipos_vigentes', [])
        total_equipos = len(equipos)
        operativos = sum(1 for equipo in equipos if equipo.estado == 'en_uso')
        en_mantenimiento = sum(1 for equipo in equipos if equipo.estado == 'en_mantenimiento')
        dañados = sum(1 for equipo in equipos if equipo.estado == 'dañado')
        local_name = (
            instance.edificio.local.nombre
            if instance.edificio and instance.edificio.local
            else "—"
        )
        building_name = instance.edificio.nombre if instance.edificio else (instance.pabellon or "—")

        builder = ExcelReportBuilder(
            title=f"Ficha Técnica de Espacio - {instance.codigo_espacio}",
            subtitle=f"{instance.get_tipo_display()} · {building_name} · Piso {instance.piso}",
            user_name=actor_name,
            filters_text=f"Sede: {local_name} | Pabellón: {building_name} | Piso: {instance.piso}",
            sheet_name="Ficha Técnica",
        )
        builder.add_kpis([
            {"label": "Total Equipos", "value": total_equipos, "tone": "info"},
            {"label": "Operativos", "value": operativos, "tone": "success"},
            {
                "label": "En Mantenimiento",
                "value": en_mantenimiento,
                "tone": "warning" if en_mantenimiento > 0 else "neutral",
            },
            {
                "label": "Con Falla / Dañados",
                "value": dañados,
                "tone": "danger" if dañados > 0 else "neutral",
            },
        ])

        assignments = getattr(instance, 'asignaciones_activas', [])
        responsible_names = [
            (
                f"{assignment.usuario.nombre} {assignment.usuario.apellido} "
                f"({assignment.get_tipo_responsabilidad_display()})"
            ).strip()
            for assignment in assignments
            if assignment.usuario
        ]
        builder.set_columns([
            {"key": "propiedad", "label": "Parámetro / Especificación", "width": 28},
            {"key": "valor", "label": "Detalle Registrado", "width": 45},
        ])
        builder.add_rows([
            {"propiedad": "Código de Espacio", "valor": instance.codigo_espacio},
            {"propiedad": "Tipo de Espacio", "valor": instance.get_tipo_display()},
            {"propiedad": "Sede / Local", "valor": local_name},
            {"propiedad": "Pabellón / Edificio", "valor": building_name},
            {"propiedad": "Piso", "valor": f"Piso {instance.piso}"},
            {
                "propiedad": "Estado Operativo del Espacio",
                "valor": "Activo" if instance.activo else "Inactivo",
            },
            {
                "propiedad": "Personal y Responsables",
                "valor": "; ".join(responsible_names) if responsible_names else "Sin personal asignado",
            },
            {"propiedad": "Total de Equipos Instalados", "valor": f"{total_equipos} unidades"},
            {
                "propiedad": "Disponibilidad Operativa",
                "valor": (
                    f"{(operativos / total_equipos * 100):.1f}%"
                    if total_equipos > 0 else "Sin equipos"
                ),
            },
        ])

        builder.add_sheet("Inventario de Equipos")
        builder.add_kpis([
            {"label": "Equipos Listados", "value": total_equipos, "tone": "info"},
            {"label": "Operativos", "value": operativos, "tone": "success"},
            {
                "label": "En Mantenimiento",
                "value": en_mantenimiento,
                "tone": "warning" if en_mantenimiento > 0 else "neutral",
            },
            {"label": "Con Falla", "value": dañados, "tone": "danger" if dañados > 0 else "neutral"},
        ])
        builder.set_columns([
            {"key": "codigo", "label": "Código Equipo", "width": 16, "align": "center"},
            {"key": "tipo_equipo", "label": "Tipo de Hardware", "width": 16},
            {"key": "marca", "label": "Marca", "width": 16},
            {"key": "modelo", "label": "Modelo", "width": 20},
            {"key": "numero_serie", "label": "N° Serie", "width": 20, "align": "center"},
            {"key": "numero_mac", "label": "Dirección MAC", "width": 20, "align": "center", "mono": True},
            {"key": "ip_actual", "label": "Dirección IP", "width": 18, "align": "center", "mono": True},
            {"key": "modo_adquisicion", "label": "Adquisición", "width": 16, "align": "center"},
            {"key": "estado", "label": "Estado", "width": 18, "align": "center"},
        ])
        builder.add_rows([
            {
                "codigo": equipo.codigo,
                "tipo_equipo": (
                    equipo.get_tipo_equipo_display()
                    if hasattr(equipo, 'get_tipo_equipo_display') else equipo.tipo_equipo
                ),
                "marca": equipo.marca,
                "modelo": equipo.modelo,
                "numero_serie": equipo.numero_serie,
                "numero_mac": equipo.numero_mac or "—",
                "ip_actual": equipo.ipv4 or equipo.ipv6 or "—",
                "modo_adquisicion": (
                    equipo.get_modo_adquisicion_display()
                    if hasattr(equipo, 'get_modo_adquisicion_display') else equipo.modo_adquisicion
                ),
                "estado": (
                    equipo.get_estado_display()
                    if hasattr(equipo, 'get_estado_display') else equipo.estado
                ),
            }
            for equipo in equipos
        ], status_keys=['estado'])

        configuration = instance.configuracion_plano or {}
        rows = configuration.get('filas', 0)
        columns = configuration.get('columnas', 0)
        seats = configuration.get('puestos', [])
        if rows > 0 and columns > 0:
            PlanoEspacioExcelReport._render_floor_plan(
                builder,
                instance,
                equipos,
                rows,
                columns,
                seats,
            )

        filename = f"ficha_plano_{instance.codigo_espacio}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return builder.to_response(filename)

    @staticmethod
    def _render_floor_plan(builder, instance, equipos, row_count, column_count, seats):
        worksheet = builder.workbook.create_sheet(title="Distribución del Plano")
        worksheet.views.sheetView[0].showGridLines = True
        worksheet.merge_cells("A1:M1")
        banner = worksheet["A1"]
        banner.value = f"  KAIROS · DISTRIBUCIÓN FÍSICA Y CROQUIS - {instance.codigo_espacio}"
        banner.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        banner.fill = PatternFill(start_color=builder.COLOR_NAVY_DARK, fill_type="solid")
        banner.alignment = Alignment(vertical="center", horizontal="left")
        worksheet.row_dimensions[1].height = 24

        worksheet.merge_cells("A2:M2")
        title = worksheet["A2"]
        title.value = f"  MAPA DE DISTRIBUCIÓN EN SALA ({row_count} Filas × {column_count} Columnas)"
        title.font = Font(name="Calibri", size=13, bold=True, color=builder.COLOR_NAVY_DARK)
        title.alignment = Alignment(vertical="center", horizontal="left")
        worksheet.row_dimensions[2].height = 26
        worksheet["A4"] = "LEYENDA:"
        worksheet["A4"].font = Font(name="Calibri", size=9, bold=True, color=builder.COLOR_TEXT_MAIN)

        legend = [
            ("B4", "Operativo (En uso)", "D1FAE5", "065F46"),
            ("D4", "En Mantenimiento", "FEF3C7", "92400E"),
            ("F4", "Con Falla / Dañado", "FEE2E2", "991B1B"),
            ("H4", "Estación Docente", "EDE9FE", "5B21B6"),
            ("J4", "Puesto Vacío", "F1F5F9", "64748B"),
        ]
        thin_border = Border(
            top=Side(style="thin", color="CBD5E1"),
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )
        for cell_reference, text, fill, font_color in legend:
            cell = worksheet[cell_reference]
            cell.value = f"  {text}  "
            cell.font = Font(name="Calibri", size=8, bold=True, color=font_color)
            cell.fill = PatternFill(start_color=fill, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        worksheet.row_dimensions[4].height = 20

        equipment_by_id = {equipo.id: equipo for equipo in equipos}
        seat_matrix = {(seat['fila'], seat['columna']): seat for seat in seats}
        start_grid_row = 6
        start_grid_column = 2
        for column_index in range(1, column_count + 1):
            column_letter = get_column_letter(start_grid_column + column_index - 1)
            worksheet.column_dimensions[column_letter].width = 16
            cell = worksheet[f"{column_letter}{start_grid_row}"]
            cell.value = f"Col {column_index}"
            cell.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color=builder.COLOR_NAVY_LIGHT, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        worksheet.row_dimensions[start_grid_row].height = 22

        grid_border = Border(
            top=Side(style="thin", color="94A3B8"),
            left=Side(style="thin", color="94A3B8"),
            right=Side(style="thin", color="94A3B8"),
            bottom=Side(style="thin", color="94A3B8"),
        )
        for row_index in range(1, row_count + 1):
            current_row = start_grid_row + row_index
            worksheet.row_dimensions[current_row].height = 36
            row_label = worksheet[f"A{current_row}"]
            row_label.value = f"Fila {row_index}"
            row_label.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
            row_label.fill = PatternFill(start_color=builder.COLOR_NAVY_LIGHT, fill_type="solid")
            row_label.alignment = Alignment(horizontal="center", vertical="center")
            row_label.border = thin_border
            worksheet.column_dimensions["A"].width = 10

            for column_index in range(1, column_count + 1):
                column_letter = get_column_letter(start_grid_column + column_index - 1)
                cell = worksheet[f"{column_letter}{current_row}"]
                cell.border = grid_border
                seat = seat_matrix.get((row_index, column_index))
                if not seat:
                    cell.value = "—"
                    cell.fill = PatternFill(start_color="F8FAFC", fill_type="solid")
                    cell.font = Font(name="Calibri", size=9, color="94A3B8")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    continue

                equipo = equipment_by_id.get(seat.get('equipo_id'))
                equipment_code = equipo.codigo if equipo else f"EQ-{seat.get('equipo_id')}"
                is_teacher_seat = seat.get('es_docente', False)
                cell.value = f"{equipment_code}{' (Docente)' if is_teacher_seat else ''}"
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                status = equipo.estado if equipo else 'en_uso'
                if is_teacher_seat:
                    fill, color = "EDE9FE", "5B21B6"
                elif status == 'en_mantenimiento':
                    fill, color = "FEF3C7", "92400E"
                elif status == 'dañado':
                    fill, color = "FEE2E2", "991B1B"
                else:
                    fill, color = "D1FAE5", "065F46"
                cell.fill = PatternFill(start_color=fill, fill_type="solid")
                cell.font = Font(name="Calibri", size=9, bold=True, color=color)
